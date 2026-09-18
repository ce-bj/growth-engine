# -*- coding: utf-8 -*-
"""一次性导入：指标字典 xlsx → config/metric_dict.json。

生成后 ``config/metric_dict.json`` 即为唯一事实源，运营/PM 直接改那份 JSON。
本脚本只留档，说明每条规则是怎么从「异常条件」那段中文翻过来的，不要再执行覆盖。

用法（仅首次）：
    python -m growth_metrics_service.tools.import_from_xlsx
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

SRC_XLSX = Path(
    r"C:\Users\范勇健\.qoderwork\workspace\mqaotdz1p13gu04b\outputs\growth-engine"
    r"\与本项目原型代码无关\outputs\attribution-flow\指标字典与假设方案库.xlsx"
)
OUT = Path(__file__).resolve().parents[1] / "config" / "metric_dict.json"

# ---------------------------------------------------------------- 规则原子简写

def pct(direction: str, gt: float) -> dict[str, Any]:
    """环比涨跌幅。"""
    return {"type": "pct_change", "dir": direction, "gt": gt}


def pp(direction: str, gt: float) -> dict[str, Any]:
    """绝对变化，单位百分点。"""
    return {"type": "abs_change_pp", "dir": direction, "gt": gt}


def pp_abs(direction: str, gt: float) -> dict[str, Any]:
    """跟上期比涨跌多少「个点」（P75、出口率占比用）。"""
    return {"type": "pp_change", "dir": direction, "gt": gt}


# 最常见的一档：「环比跌 >10% 且绝对跌 >2pp（须同时满足）」
RATE_DOWN = {"all": [pct("down", 0.10), pp("down", 2)]}
RATE_UP = {"all": [pct("up", 0.10), pp("up", 2)]}
DEFAULT_JUDGE = {"anomaly": RATE_DOWN, "improve": RATE_UP}

# 分母人数 <10 不计算
BASE_DENOM10 = [{"type": "denominator", "gte": 10}]
# 周期 <14 天不计算
BASE_PERIOD14 = {"type": "period_days", "gte": 14}
# 该渠道 UV <30 不计算
BASE_CH_UV30 = [{"type": "companion", "ref": "UV@渠道", "gte": 30}]
# 该词本期点击 <10 不计算
BASE_KW_CLICK10 = [{"type": "companion", "ref": "来路关键词点击量@搜索词", "gte": 10}]

# ------------------------------------------------------- 逐条覆盖（按 xlsx 序号）
# 未列出的条目走 DEFAULT_JUDGE + role=primary + tri_state=full。

OVERRIDES: dict[int, dict[str, Any]] = {
    1: {  # 留资数@全站
        "min_base": [{"type": "prev_value", "gt": 5}],
        "judge": {"anomaly": pct("down", 0.30), "improve": pct("up", 0.30)},
        "deadband_ref": "留资数@全站",
    },
    2: {  # 留资率@全站 · 北极星
        "min_base": [],
        "judge": {
            "anomaly": {"any": [pct("down", 0.20), {"type": "value", "lt": 0.005}]},
            "improve": {"any": [pct("up", 0.20), {"type": "value", "gt": 0.025}]},
        },
        "note": "另比年度目标 3%",
    },
    3: {"min_base": [BASE_PERIOD14]},  # 有效浏览率@全站 · 流失点1 主指标
    4: {"min_base": BASE_DENOM10 + [BASE_PERIOD14]},  # 转化交互率@全站
    5: {"min_base": BASE_DENOM10 + [BASE_PERIOD14]},  # 留资完成率@全站
    6: {  # UV@全站
        "min_base": [{"type": "prev_value", "gt": 100}],
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
    },
    7: {  # 渠道UV占比@渠道
        "tri_state": "no_flat",
        "min_base": BASE_CH_UV30,
        "judge": {
            "anomaly": {
                "all": [
                    pct("up", 0.30),
                    {"type": "ref_ratio", "left": "留资率@渠道", "right": "留资率@全站", "lt": 0.5},
                ],
            },
            "improve": {
                "all": [
                    pct("up", 0.30),
                    {"type": "ref_ratio", "left": "留资率@渠道", "right": "留资率@全站", "gte": 1.0},
                ],
            },
        },
    },
    8: {  # 秒退率@全站 · 升是坏
        "min_base": [BASE_PERIOD14],
        "judge": {"anomaly": pct("up", 0.15), "improve": pct("down", 0.15)},
    },
    9: {"tri_state": "none", "role": "counting", "min_base": []},   # 有效浏览访客
    10: {"tri_state": "none", "role": "counting", "min_base": []},  # 转化交互访客
    11: {"tri_state": "none", "role": "counting", "min_base": []},  # 留资访客
    12: {  # 健康度待处理问题数
        "tri_state": "anomaly_only",
        "min_base": [],
        "judge": {"anomaly": {"type": "health_high_risk"}},
    },
    13: {  # 健康度待处理问题列表
        "tri_state": "anomaly_only",
        "role": "reference",
        "min_base": [],
        "judge": {"anomaly": {"type": "health_high_risk"}},
    },
    14: {  # UV@渠道
        "min_base": [{"type": "self_value", "gte": 30}],
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
    },
    15: {  # 留资率@渠道
        "min_base": BASE_CH_UV30,
        "judge": {
            "anomaly": {
                "any": [
                    pct("down", 0.20),
                    {"type": "ref_ratio", "left": "留资率@渠道", "right": "留资率@全站", "lt": 0.5},
                ],
            },
            "improve": pct("up", 0.20),
        },
    },
    16: {  # 留资数@渠道
        "min_base": [{"type": "prev_value", "gt": 5}],
        "judge": {"anomaly": pct("down", 0.30), "improve": pct("up", 0.30)},
    },
    17: {  # 来路关键词展现量
        "min_base": BASE_KW_CLICK10,
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
    },
    18: {  # 来路关键词点击量
        "min_base": [{"type": "self_value", "gte": 10}],
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
    },
    19: {"min_base": BASE_CH_UV30},  # 有效浏览率@渠道
    20: {  # 秒退率@渠道
        "min_base": BASE_CH_UV30,
        "judge": {"anomaly": pct("up", 0.15), "improve": pct("down", 0.15)},
    },
    21: {  # 落地接住率@页面
        "min_base": BASE_DENOM10,
        "note": "三个通过率都没掉时仍看本条或落地停留，有页破线仍定位流失点 1",
    },
    22: {"min_base": BASE_DENOM10},  # 合格浏览率@页面
    23: {"min_base": BASE_DENOM10},  # 落地接住率@渠道×页
    24: {"min_base": BASE_DENOM10},  # 合格浏览率@渠道×页
    25: {  # 落地接住率@搜索词×页
        "min_base": BASE_KW_CLICK10,
        "judge": {
            "anomaly": {
                "any": [
                    RATE_DOWN,
                    {
                        "type": "companion_pct_change",
                        "ref": "来路关键词点击量@搜索词",
                        "dir": "down",
                        "gt": 0.20,
                    },
                ],
            },
            "improve": RATE_UP,
        },
    },
    26: {  # 落地停留均值@页面 · 旁证
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
        "label_on_anomaly": "停留明显变短",
    },
    27: {
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
        "label_on_anomaly": "停留明显变短",
    },
    28: {  # 滚动深度P75@页面 · 旁证
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pp_abs("down", 15), "improve": pp_abs("up", 15)},
        "label_rules": [{"when": {"type": "value", "lt": 0.40}, "label": "多数人还在首屏附近"}],
    },
    29: {  # 按词页面停留 · 不判三态
        "tri_state": "none",
        "role": "corroborating",
        "min_base": BASE_KW_CLICK10,
    },
    30: {"min_base": BASE_DENOM10},  # 转化交互率@渠道
    31: {"min_base": BASE_DENOM10},
    32: {"min_base": BASE_DENOM10},
    33: {"min_base": BASE_DENOM10},
    34: {"min_base": BASE_DENOM10},
    35: {"min_base": BASE_DENOM10},
    36: {"min_base": BASE_DENOM10},
    37: {  # 未转化页出口率@页面 · 升是坏，按占比个点
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pp_abs("up", 10), "improve": pp_abs("down", 10)},
    },
    38: {"min_base": BASE_DENOM10},  # 转化交互率@页面
    39: {"min_base": BASE_DENOM10},  # 转化交互率@渠道×页
    40: {  # 人均有效浏览页数 · 旁证
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
        "label_on_anomaly": "逛得更浅",
    },
    41: {
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("down", 0.20), "improve": pct("up", 0.20)},
    },
    42: {"tri_state": "none", "role": "corroborating", "min_base": BASE_DENOM10},
    43: {"tri_state": "none", "role": "reference", "min_base": []},
    44: {"tri_state": "none", "role": "reference", "min_base": []},
    45: {  # 站内搜索0结果率 · 升是坏
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("up", 0.15), "improve": pct("down", 0.15)},
    },
    46: {  # 搜后离站率 · 升是坏
        "role": "corroborating",
        "min_base": BASE_DENOM10,
        "judge": {"anomaly": pct("up", 0.15), "improve": pct("down", 0.15)},
    },
    47: {"tri_state": "none", "role": "reference", "min_base": []},
    48: {  # 表单曝光率 · 未接入
        "tri_state": "none",
        "role": "reference",
        "min_base": [{"type": "unavailable", "status": "not_connected", "reason": "埋点未接入"}],
    },
    49: {  # 页面内容 · 读页
        "tri_state": "none",
        "role": "reference",
        "min_base": [{"type": "page_fetch"}],
    },
    50: {  # 询价入口个数 · 读页，供假设 9 用 value==0
        "tri_state": "none",
        "role": "reference",
        "min_base": [{"type": "page_fetch"}],
    },
    51: {"min_base": BASE_DENOM10},  # 留资完成率@渠道
    52: {"min_base": BASE_DENOM10},  # 表单完成率@全站
    53: {"min_base": BASE_DENOM10},  # 表单完成率@渠道
    54: {"min_base": BASE_DENOM10},  # 开聊留资率@全站
    55: {"min_base": BASE_DENOM10},  # 开聊留资率@渠道
    56: {"tri_state": "none", "role": "reference", "min_base": []},
    57: {"tri_state": "none", "role": "corroborating", "min_base": []},
    58: {  # 高价值线索占比 · 一期不做
        "tri_state": "none",
        "role": "reference",
        "min_base": [{"type": "unavailable", "status": "not_in_phase1", "reason": "一期不进本层"}],
    },
    59: {  # 页面留资率 · 现网无
        "tri_state": "none",
        "role": "reference",
        "min_base": [{"type": "unavailable", "status": "not_connected", "reason": "现网无，二期留口"}],
    },
}


def build() -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(SRC_XLSX, read_only=True, data_only=True)
    ws = wb["指标字典"]
    rows = [
        ["" if c is None else str(c).strip() for c in row]
        for row in ws.iter_rows(values_only=True)
    ]
    metrics: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        seq = int(row[0])
        name, numer, denom, meaning, dim, layer, min_base_raw, tri_raw, anomaly_raw = row[1:10]
        ov = OVERRIDES.get(seq, {})
        entry: dict[str, Any] = {
            "id": f"M{seq}",
            "ref": f"{name}@{dim}",
            "name": name,
            "dim": dim,
            "layer": layer,
            "role": ov.get("role", "primary"),
            "tri_state": ov.get("tri_state", "full"),
            "numerator": numer,
            "denominator": denom,
            "meaning": meaning,
            "min_base": {"raw": min_base_raw, "rules": ov.get("min_base", BASE_DENOM10)},
            "judge": ov.get("judge", DEFAULT_JUDGE) if ov.get("tri_state", "full") != "none" else None,
            "baseline": "past_4w_same_period",
            "source_raw": {"是否判3态": tri_raw, "异常条件": anomaly_raw},
        }
        for extra in ("note", "label_on_anomaly", "label_rules", "deadband_ref"):
            if extra in ov:
                entry[extra] = ov[extra]
        metrics.append(entry)

    refs = [m["ref"] for m in metrics]
    dupes = {r for r in refs if refs.count(r) > 1}
    if dupes:
        raise SystemExit(f"ref 冲突：{dupes}")

    return {
        "version": "1.0.0",
        "source": "指标字典与假设方案库.xlsx · 指标字典",
        "note": "生成后此文件为唯一事实源，请直接编辑，不要再从 xlsx 覆盖生成",
        "dims": ["全站", "渠道", "页面", "渠道×页", "搜索词", "搜索词×页", "通道", "路径"],
        "layers": ["扫描", "访问入口", "流失点1", "流失点2", "流失点3", "技术侧"],
        "metrics": metrics,
    }


if __name__ == "__main__":
    data = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {OUT} metrics={len(data['metrics'])}")
